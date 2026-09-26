#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""PHASE 27 — bilingual Persian-content inventory Excel.
Sheets:
  1. Persian Content  — every CODE occurrence (non-comment): exact phrase,
                        folder, file, line, type, scope, English translation
  2. Unique Strings   — unique non-comment strings + translations + counts
  3. Comments         — every comment occurrence (complete inventory)
  4. Summary          — stats + mechanism notes
Output: download/i18n-persian-inventory.xlsx
"""

import json
import os
import sys

XLSX_SKILL_DIR = "/home/z/my-project/skills/xlsx"
for sub in [XLSX_SKILL_DIR, os.path.join(XLSX_SKILL_DIR, "templates")]:
    if sub not in sys.path:
        sys.path.insert(0, sub)

from base import (  # noqa: E402
    FONT_NAME, PRIMARY, NEUTRAL_600, NEUTRAL_900,
    font_body, font_caption,
    setup_sheet, style_header_row, style_data_row,
)
from openpyxl import Workbook  # noqa: E402
from openpyxl.styles import Alignment, Font  # noqa: E402
from openpyxl.utils import get_column_letter  # noqa: E402

DIR = "/home/z/my-project/scripts/i18n"
OUT = "/home/z/my-project/download/i18n-persian-inventory.xlsx"

occ = json.load(open(os.path.join(DIR, "occurrences.json"), encoding="utf-8"))
translations = json.load(open(os.path.join(DIR, "translations.json"), encoding="utf-8"))
extras = json.load(open(os.path.join(DIR, "extra-translations.json"), encoding="utf-8"))
comment_tr_path = os.path.join(DIR, "comment-translations.json")
comment_tr = json.load(open(comment_tr_path, encoding="utf-8")) if os.path.exists(comment_tr_path) else {}
stats = json.load(open(os.path.join(DIR, "stats.json"), encoding="utf-8"))

TR = {**translations, **extras}

def scope_of(rel):
    return "server" if (rel.startswith("app/api/") or rel == "proxy.ts") else "client"

def tr_for(key):
    v = TR.get(key)
    return v if isinstance(v, str) else ""

def split_path(rel):
    # rel is like components/modules/finance/payroll-page.tsx
    parts = rel.split("/")
    folder = "/".join(parts[:-1]) or "(root)"
    fname = parts[-1]
    return folder, fname

TYPE_FA = {
    "jsx_text": "متن JSX",
    "toast": "توست",
    "error": "خطا",
    "template_param": "تمپلیت پارامتری",
    "template": "تمپلیت",
    "literal": "لیترال",
    "comment": "کامنت",
    "console": "کنسول",
}
def type_fa(t):
    if ":" in t:
        kind, name = t.split(":", 1)
        fa = TYPE_FA.get(kind, kind)
        return f"{fa} ({name})"
    return TYPE_FA.get(t, t)

wb = Workbook()

# ─── Sheet 1: Persian Content (per-occurrence, code only) ────────────────
ws = wb.active
ws.title = "Persian Content"
headers = ["#", "عبارت دقیق فارسی (Persian)", "ترجمه انگلیسی (English)",
           "پوشه (Folder)", "فایل (File)", "خط (Line)",
           "نوع (Type)", "دامنه (Scope)"]
last_col = len(headers) + 1
setup_sheet(ws, title="Printoo24 ERP — فهرست کامل محتوای فارسی (کد) / Persian UI Content Inventory", last_col=last_col)
for col_idx, h in enumerate(headers, start=2):
    ws.cell(row=4, column=col_idx, value=h)
style_header_row(ws, row_num=4, col_start=2, col_end=last_col)

code_occ = [o for o in occ if o["type"] != "comment"]
row = 5
for i, o in enumerate(code_occ, 1):
    folder, fname = split_path(o["file"])
    key = o.get("key") or o["text"]
    vals = [i, o["text"], tr_for(key), folder, fname, o["line"],
            type_fa(o["type"]), scope_of(o["file"])]
    for col_idx, v in enumerate(vals, start=2):
        ws.cell(row=row, column=col_idx, value=v)
    style_data_row(ws, row_num=row, col_start=2, col_end=last_col, row_index=i - 1)
    row += 1
n_code = row - 5

widths = {2: 7, 3: 42, 4: 42, 5: 30, 6: 22, 7: 8, 8: 20, 9: 10}
for c, w in widths.items():
    ws.column_dimensions[get_column_letter(c)].width = w
ws.freeze_panes = "C5"
ws.auto_filter.ref = f"B4:{get_column_letter(last_col)}{row - 1}"

# ─── Sheet 2: Unique Strings ─────────────────────────────────────────────
ws2 = wb.create_sheet("Unique Strings")
headers2 = ["#", "عبارت فارسی (Persian)", "ترجمه انگلیسی (English)",
            "تعداد تکرار (Count)", "نوع‌ها (Types)", "کلاینت (Client)", "سرور (Server)"]
last_col2 = len(headers2) + 1
setup_sheet(ws2, title="رشته‌های یکتا + ترجمه / Unique Strings + Translations", last_col=last_col2)
for col_idx, h in enumerate(headers2, start=2):
    ws2.cell(row=4, column=col_idx, value=h)
style_header_row(ws2, row_num=4, col_start=2, col_end=last_col2)

comment_keys = {o.get("key") or o["text"] for o in occ if o["type"] == "comment"}
catalog = json.load(open(os.path.join(DIR, "catalog.json"), encoding="utf-8"))["strings"]
uniq = [(k, v) for k, v in catalog.items() if k not in comment_keys]
uniq.sort(key=lambda kv: -kv[1]["count"])
row = 5
missing = 0
for i, (k, v) in enumerate(uniq, 1):
    trn = tr_for(k)
    if not trn:
        missing += 1
    types = ", ".join(f"{t}×{n}" for t, n in sorted(v["types"].items(), key=lambda x: -x[1])[:3])
    vals = [i, k, trn, v["count"], types,
            "✓" if v["client_files"] else "—", "✓" if v["server_files"] else "—"]
    for col_idx, val in enumerate(vals, start=2):
        ws2.cell(row=row, column=col_idx, value=val)
    style_data_row(ws2, row_num=row, col_start=2, col_end=last_col2, row_index=i - 1)
    row += 1
for c, w in {2: 7, 3: 48, 4: 48, 5: 12, 6: 26, 7: 10, 8: 10}.items():
    ws2.column_dimensions[get_column_letter(c)].width = w
ws2.freeze_panes = "C5"
ws2.auto_filter.ref = f"B4:{get_column_letter(last_col2)}{row - 1}"

# ─── Sheet 3: Comments ───────────────────────────────────────────────────
ws3 = wb.create_sheet("Comments")
headers3 = ["#", "کامنت فارسی (Persian)", "ترجمه انگلیسی (English)",
            "پوشه (Folder)", "فایل (File)", "خط (Line)"]
last_col3 = len(headers3) + 1
setup_sheet(ws3, title="کامنت‌های فارسی / Persian Code Comments", last_col=last_col3)
for col_idx, h in enumerate(headers3, start=2):
    ws3.cell(row=4, column=col_idx, value=h)
style_header_row(ws3, row_num=4, col_start=2, col_end=last_col3)
row = 5
comment_occ = [o for o in occ if o["type"] == "comment"]
for i, o in enumerate(comment_occ, 1):
    folder, fname = split_path(o["file"])
    key = o.get("key") or o["text"]
    vals = [i, o["text"], comment_tr.get(key, ""), folder, fname, o["line"]]
    for col_idx, v in enumerate(vals, start=2):
        ws3.cell(row=row, column=col_idx, value=v)
    if i % 2 == 1:
        style_data_row(ws3, row_num=row, col_start=2, col_end=last_col3, row_index=i - 1)
    row += 1
for c, w in {2: 7, 3: 60, 4: 60, 5: 28, 6: 20, 7: 8}.items():
    ws3.column_dimensions[get_column_letter(c)].width = w
ws3.freeze_panes = "C5"
ws3.auto_filter.ref = f"B4:{get_column_letter(last_col3)}{row - 1}"

# ─── Sheet 4: Summary ────────────────────────────────────────────────────
ws4 = wb.create_sheet("Summary")
setup_sheet(ws4, title="خلاصه / Summary — Phase 27 Bilingual Migration", last_col=4)
rows4 = [
    ("فایل‌های اسکن‌شده / Files scanned", stats["files_scanned"]),
    ("کل رخدادهای فارسی / Total Persian occurrences", stats["total_occurrences"]),
    ("— در فایل‌های کلاینت / client occurrences", stats["client_occurrences"]),
    ("— در فایل‌های سرور / server occurrences", stats["server_occurrences"]),
    ("رخدادهای کد (غیر کامنت) / Code occurrences (non-comment)", len(code_occ)),
    ("رخدادهای کامنت / Comment occurrences", len(comment_occ)),
    ("رشته‌های یکتای کد / Unique code strings", len(uniq)),
    ("رشته‌های یکتای کامنت / Unique comment strings", len(comment_keys)),
    ("ترجمه‌شده (کد) / Translated (code)", len(uniq) - missing),
    ("ترجمه‌شده (کامنت) / Translated (comments)", len(comment_tr)),
    ("", ""),
    ("مکانیزم / Mechanism", "t() از @/lib/i18n — کلید طبیعی، دیفالت EN، سوییچ per-profile"),
    ("مستندات اجباری / Mandatory docs", "I18N_RULES.md در ریشهٔ ریپو"),
    ("ابزارها / Tooling", "scripts/i18n/: extract → translate → build-dicts"),
]
r = 4
for label, val in rows4:
    ws4.cell(row=r, column=2, value=label).font = font_body()
    c = ws4.cell(row=r, column=3, value=val)
    c.font = Font(name=FONT_NAME, size=11, color=PRIMARY, bold=False)
    r += 1
ws4.column_dimensions["B"].width = 52
ws4.column_dimensions["C"].width = 58
note = ws4.cell(row=r + 1, column=2,
    value="یادداشت: رشته‌های سمت سرور در API فارسی می‌مانند و کلاینت در لحظهٔ نمایش ترجمه می‌کند (choke-point: lib/api.ts). "
          "کلیدهای طبیعی یعنی t() حتی مقادیر داینامیک دیتابیس را هم ترجمه می‌کند.")
note.font = font_caption()
note.alignment = Alignment(horizontal="left", vertical="top", wrap_text=True)
ws4.merge_cells(start_row=r + 1, start_column=2, end_row=r + 3, end_column=4)

wb.properties.creator = "Z.ai"
os.makedirs(os.path.dirname(OUT), exist_ok=True)
wb.save(OUT)
print(f"saved: {OUT}")
print(f"sheet1 rows: {n_code}, sheet2 unique: {len(uniq)} (missing translation: {missing}), sheet3 comments: {len(comment_occ)} (translated: {len(comment_tr)})")
